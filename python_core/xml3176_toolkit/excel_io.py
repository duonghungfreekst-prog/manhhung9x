from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .models import Xml3176Document, XmlSheet
from .schema import XML_STRUCTURE
from .parser import META_COLS


class Xml3176ExcelIO:
    """Xuất/nhập Excel theo từng sheet XML.

    v0.3 bỏ sheet _META. Việc ghép hồ sơ khi Excel -> XML dựa trên MA_LK.
    Các cột nội bộ __... không xuất ra Excel để người dùng không bị rối.
    """

    @staticmethod
    def export_excel(doc: Xml3176Document, path: str | Path, hoso_indexes: Optional[List[int]] = None, ma_lks: Optional[List[str]] = None) -> None:
        from .parser import Xml3176Parser
        if hoso_indexes is not None or ma_lks is not None:
            doc = Xml3176Parser.filter_doc(doc, hoso_indexes=hoso_indexes, ma_lks=ma_lks)
        wb = Workbook()
        first = True
        for loai in doc.sheet_names():
            sheet = doc.sheets[loai]
            sws = wb.active if first else wb.create_sheet(loai)
            first = False
            sws.title = loai
            columns = [c for c in sheet.columns if not c.startswith("__")]
            # MA_LK phải đứng đầu nếu có để người dùng nhìn và quản lý liên kết.
            if "MA_LK" in columns:
                columns = ["MA_LK"] + [c for c in columns if c != "MA_LK"]
            sws.append(columns)
            for row in sheet.rows:
                sws.append([row.get(c, "") for c in columns])
            Xml3176ExcelIO._style_table(sws)
        if not doc.sheet_names():
            ws = wb.active
            ws.title = "XML1"
            ws.append(["MA_LK"])
        wb.save(path)

    @staticmethod
    def import_excel(path: str | Path, macskcb: str = "", ngaylap: str = "") -> Xml3176Document:
        wb = load_workbook(path)
        doc = Xml3176Document(macskcb=macskcb, ngaylap=ngaylap or datetime.now().strftime("%Y%m%d"), source_path=str(path))
        # MA_LK -> HOSO_INDEX. Không cần sheet Meta.
        ma_lk_to_index: Dict[str, int] = {}
        next_index = 1

        for sname in wb.sheetnames:
            loai = sname.upper().strip()
            if not loai.startswith("XML"):
                continue
            ws = wb[sname]
            columns = [str(c.value).strip() for c in ws[1] if c.value is not None and str(c.value).strip()]
            data_columns = [c for c in columns if not c.startswith("__")]
            internal_columns = ["__HOSO_INDEX"] + data_columns
            rows = []
            for cells in ws.iter_rows(min_row=2, max_col=len(columns), values_only=True):
                if cells is None:
                    continue
                row: Dict[str, str] = {}
                for i, col in enumerate(columns):
                    value = cells[i] if i < len(cells) else None
                    row[col] = "" if value is None else str(value)
                if not any(v != "" for k, v in row.items() if not k.startswith("__")):
                    continue
                ma_lk = str(row.get("MA_LK", "") or "").strip()
                if ma_lk:
                    if ma_lk not in ma_lk_to_index:
                        ma_lk_to_index[ma_lk] = next_index
                        next_index += 1
                    row["__HOSO_INDEX"] = str(ma_lk_to_index[ma_lk])
                else:
                    row["__HOSO_INDEX"] = str(next_index)
                    next_index += 1
                rows.append(row)
                if not doc.macskcb:
                    doc.macskcb = row.get("MA_CSKCB", "") or row.get("MACSKCB", "") or doc.macskcb
            spec = XML_STRUCTURE.get(loai, {"root": loai, "list": None, "record": loai})
            doc.sheets[loai] = XmlSheet(
                loai_hoso=loai,
                root_tag=spec["root"],
                list_tag=spec.get("list"),
                record_tag=spec["record"],
                columns=internal_columns,
                rows=rows,
            )
        return doc

    @staticmethod
    def _style_header(ws, row_num: int) -> None:
        fill = PatternFill("solid", fgColor="D9EAF7")
        thin = Side(style="thin", color="B7B7B7")
        for cell in ws[row_num]:
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    @staticmethod
    def _style_table(ws) -> None:
        if ws.max_row >= 1:
            Xml3176ExcelIO._style_header(ws, 1)
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
        thin = Side(style="thin", color="D0D0D0")
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=False)
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for col_idx in range(1, ws.max_column + 1):
            letter = get_column_letter(col_idx)
            max_len = 10
            for cell in ws[letter][:80]:
                if cell.value is not None:
                    max_len = max(max_len, min(len(str(cell.value)) + 2, 55))
            ws.column_dimensions[letter].width = max_len
